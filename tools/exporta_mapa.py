# -*- coding: utf-8 -*-
"""Gera o mapa da Controladoria do mês a partir do mapa do mês anterior.

Por que TEMPLATE e não arquivo novo: o mapa tem 6 abas, 2.398 fórmulas,
autofiltro, painéis congelados e uma formatação que ninguém quer redigitar. A
biblioteca que o hub usa no navegador (SheetJS community) lê estilo mas não
escreve — medido: a célula volta com `patternType: none`. Reconstruir do zero
daria um arquivo parecido, nunca igual.

Abrindo o mapa do mês anterior e trocando só os valores, formatação e fórmulas
não são imitadas: são as mesmas, porque é o mesmo arquivo. O custo é ~40s de
round-trip (a aba PLANTEL carrega 1.017.875 linhas de formatação vazia), o que
para um fechamento mensal não é problema.

Os números vêm do MOTOR do hub, via tools/motor_mapa.js — não de uma segunda
apuração escrita aqui. Ver o cabeçalho daquele arquivo para o porquê.

Uso:
    python tools/exporta_mapa.py 08/2026
    python tools/exporta_mapa.py 08/2026 --saida "C:/caminho/arquivo.xlsx"
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import unicodedata
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
MAPA_DIR = Path(
    "G:/Drives compartilhados/Luxor Controladoria/Relatórios Gerenciais/"
    "RELATORIOS - OPERAÇÃO HARAS E FAZENDA PG/Posição Equinos/PLANTEL - Movimentações"
)
MES_PT = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago",
          "Set", "Out", "Nov", "Dez"]
# a linha agregada das receptoras: o número é a contagem do rebanho, não nome
RX_RECEPTORAS = re.compile(r"^RECEPTORAS(\s+\d+)?$")


def norm(s) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().upper()


def motor(mes: str) -> dict:
    """Roda o motor do hub e devolve a apuração do ano até `mes`."""
    print(f"  apurando {mes} pelo motor do hub...", flush=True)
    saida = subprocess.run(
        ["node", str(RAIZ / "tools" / "motor_mapa.js"), mes],
        capture_output=True, cwd=RAIZ,
    )
    if saida.returncode != 0:
        sys.exit("motor_mapa.js falhou:\n" + saida.stderr.decode("utf-8", "replace"))
    return json.loads(saida.stdout.decode("utf-8"))


def acha_template(mes_ant: str) -> Path:
    """Mapa do mês anterior — o retrato mais recente já fechado."""
    ano, mm = mes_ant.split("-")
    rot = f"({MES_PT[int(mm) - 1]} {ano})"
    pasta = MAPA_DIR / ano
    cands = [p for p in pasta.glob("Plantel Haras Pao Grande*.xlsx")
             if not p.name.startswith("~$") and rot in p.name]
    if not cands:
        sys.exit(f"não achei o mapa de {rot} em {pasta}")
    return max(cands, key=lambda p: p.stat().st_mtime)


def cabecalho(ws, rotulo: str, ate_linha: int = 8):
    """Linha e mapa {rótulo normalizado: coluna} do cabeçalho da aba."""
    for r in range(1, ate_linha + 1):
        for c in range(1, ws.max_column + 1):
            if norm(ws.cell(row=r, column=c).value) == norm(rotulo):
                cols = {}
                for cc in range(1, ws.max_column + 1):
                    v = norm(ws.cell(row=r, column=cc).value)
                    if v:
                        cols.setdefault(v, cc)
                return r, cols
    sys.exit(f"não achei o cabeçalho '{rotulo}' na aba {ws.title}")


def copia_estilo(ws, de: int, para: int, ncols: int):
    """Uma linha nova nasce com o estilo da linha de cima — inclusive formato de
    número e bordas. Sem isso ela sai branca no meio da tabela."""
    for c in range(1, ncols + 1):
        origem, destino = ws.cell(row=de, column=c), ws.cell(row=para, column=c)
        if origem.has_style:
            destino._style = copy(origem._style)


def casa_linhas(itens, existentes, base_da_linha=None):
    """Casa cada animal com a linha que ele já tem na planilha.

    1. pelo nome de hoje;
    2. por QUALQUER nome que o animal já teve no ano. A Controladoria não
       renomeia junto com o haras: o mapa de junho ainda chama "RECEPTORAS 121"
       o rebanho que julho chama "RECEPTORAS 120", e "MORENA L2 X DAMASCO..." a
       potra que virou "POTRA MORENA L2 X DAMASCO...";
    3. pelo VALOR EM DEZ/2025, quando o renome é anterior a 2026 e portanto
       invisível para o motor — "PALADINO LUEKIM DA PAO GRANDE (CARLA)" no mapa é
       "PALADINO FIGUEIRA DA PAO GRANDE (CARLA)" no arquivo do haras, e nenhum
       log de 2026 conta isso. Aqui exige-se base idêntica E o mesmo começo de
       nome, porque base igual sozinha casaria irmãos de mesmo valor.

    Sem esses três passos, cada renome deixa DUAS linhas no mapa: a antiga com os
    valores velhos do template e uma nova com o movimento do ano."""
    linha_de, usadas, novos = {}, set(), []
    for passo in (0, 1):
        for a in itens:
            if a["chave"] in linha_de:
                continue
            nomes = [a["nome"]] if passo == 0 else a.get("aliases", [])
            for nome in nomes:
                r = existentes.get(norm(nome))
                if r and r not in usadas:
                    linha_de[a["chave"]] = r
                    usadas.add(r)
                    break
    if base_da_linha:
        livres = [(n, r) for n, r in existentes.items() if r not in usadas]
        for a in itens:
            if a["chave"] in linha_de or not a.get("base"):
                continue
            ini = norm(a["nome"])[:6]
            for n, r in livres:
                if r in usadas or not n.startswith(ini):
                    continue
                if abs(base_da_linha.get(r, 0) - a["base"]) < 0.01:
                    linha_de[a["chave"]] = r
                    usadas.add(r)
                    break
    for a in itens:
        if a["chave"] not in linha_de:
            novos.append(a)
    return linha_de, novos


def linha_totais(ws, ncols: int, desde: int):
    """Linha de TOTAIS DAS MOVIMENTAÇÕES / fim do bloco de dados."""
    for r in range(desde, min(ws.max_row, desde + 4000) + 1):
        for c in range(1, ncols + 1):
            if "TOTAIS" in norm(ws.cell(row=r, column=c).value):
                return r
    return None


def escreve_plantel(ws, dados: dict):
    """Troca o plantel pelo do mês. Colunas com fórmula (letra, idade, safra,
    PLANTEL HPG/EDUARDO) não são tocadas: elas se recalculam do que a gente
    escreve ao lado."""
    lin_cab, cols = cabecalho(ws, "NOME")
    ncols = max(cols.values())
    FORMULA = {"LETRA", "IDADE (ANO)", "SAFRA", "PLANTEL HPG", "PLANTEL EDUARDO"}

    ult = lin_cab
    for r in range(lin_cab + 1, ws.max_row + 1):
        if norm(ws.cell(row=r, column=cols["NOME"]).value):
            ult = r
        elif r - ult > 30:            # fim da tabela: o resto é formatação vazia
            break

    existentes = {}
    for r in range(lin_cab + 1, ult + 1):
        n = norm(ws.cell(row=r, column=cols["NOME"]).value)
        if n:
            existentes.setdefault(n, r)

    itens = [dict(l, chave=l["__nome"], nome=l["__nome"],
                  aliases=l.get("__aliases", [])) for l in dados["plantel"]]
    linha_de, novos = casa_linhas(itens, existentes)
    if novos:
        ws.insert_rows(ult + 1, amount=len(novos))
        for i, a in enumerate(novos):
            copia_estilo(ws, ult, ult + 1 + i, ncols)
            linha_de[a["chave"]] = ult + 1 + i

    for l in dados["plantel"]:
        r = linha_de[l["__nome"]]
        for rot, c in cols.items():
            if rot in FORMULA:
                continue
            orig = next((k for k in l if norm(k) == rot), None)
            if orig is None:
                continue
            v = l[orig]
            if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
                v = datetime.strptime(v, "%Y-%m-%d")
            ws.cell(row=r, column=c).value = v if v != "" else None
        if "COMISSAO (R$)" in cols:
            ws.cell(row=r, column=cols["COMISSAO (R$)"]).value = l["__comissao"] or None

    # linha que sobrou do mês anterior e não está mais no plantel: zera cota e
    # valor, que é como o haras marca quem saiu — apagar perderia o histórico
    vivas = set(linha_de.values())
    for n, r in existentes.items():
        if r not in vivas:
            ws.cell(row=r, column=cols["COTAS (%)"]).value = 0
            ws.cell(row=r, column=cols["VALOR (R$)"]).value = 0

    n = len(novos)
    ult_dados = ult + n
    # os totais desta aba são SUBTOTAL, e as somas de cota/valor/comissão ficam
    # mais abaixo — todas precisam alcançar as linhas novas
    for r in range(ult_dados + 1, ult_dados + 12):
        for c in range(1, ncols + 1):
            f = ws.cell(row=r, column=c).value
            if not isinstance(f, str) or not f.startswith("="):
                continue
            letra = get_column_letter(c)
            if "SUBTOTAL" in f:
                ws.cell(row=r, column=c).value = (
                    f"=SUBTOTAL(9,{letra}{lin_cab + 1}:{letra}{ult_dados})")
            elif f.startswith("=SUM("):
                ws.cell(row=r, column=c).value = (
                    f"=SUM({letra}{lin_cab + 1}:{letra}{ult_dados})")
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A{lin_cab}:{get_column_letter(ncols)}{ult_dados}"
    return len(dados["plantel"]), n


def escreve_movimentacoes(ws, dados: dict, rot_mes: str, desloca_plantel: int):
    """A aba é CUMULATIVA (jan a dez na mesma linha), então o que entra é o
    acumulado do ano. Linha que já existe é atualizada no lugar; animal novo
    entra antes da linha de totais.

    Os totais são SUBTOTAL(9,...), não SOMA: eles respeitam o autofiltro, que é
    como a Controladoria olha só a Carla sem mexer em fórmula nenhuma. Trocar
    por SUM mudaria o comportamento da planilha — o que muda aqui é o intervalo,
    não a função."""
    lin_cab, cols = cabecalho(ws, "NOME")
    col = {
        "nome": cols["NOME"], "sufixo": cols["SUFIXO"], "categoria": cols["CATEGORIA"],
        "status": cols["STATUS PLANTEL"], "cota": cols["PROPRIETARIO - HPG"],
        "base": next(v for k, v in cols.items() if k.startswith("VALOR EM DEZ")),
        "compras": cols["COMPRAS"],
        "embrioes": cols["EMBRIOES CONFIRMADOS ACIMA DE 60 DIAS"],
        "venda": cols["BAIXAS POR VENDA"],
        "morte_doacao": cols["BAIXAS POR MORTE/DOACAO"],
        "reavaliacao": cols["REAVALIACAO"],
        "final": next(v for k, v in cols.items()
                      if k.startswith("VALOR EM") and not k.startswith("VALOR EM DEZ")),
    }
    ncols = max(cols.values())
    lin_tot = linha_totais(ws, ncols, lin_cab + 1)
    if not lin_tot:
        sys.exit("não achei a linha de TOTAIS DAS MOVIMENTAÇÕES")

    existentes = {}
    for r in range(lin_cab + 1, lin_tot):
        n = norm(ws.cell(row=r, column=col["nome"]).value)
        if n:
            existentes.setdefault(n, r)

    base_da_linha = {}
    for r in range(lin_cab + 1, lin_tot):
        v = ws.cell(row=r, column=col["base"]).value
        base_da_linha[r] = float(v) if isinstance(v, (int, float)) else 0.0
    linha_de, novos = casa_linhas(dados["movimentacoes"], existentes, base_da_linha)
    if novos:
        ws.insert_rows(lin_tot, amount=len(novos))
        for i, a in enumerate(novos):
            r = lin_tot + i
            copia_estilo(ws, lin_tot - 1, r, ncols)
            linha_de[a["chave"]] = r
        lin_tot += len(novos)

    for a in dados["movimentacoes"]:
        r = linha_de[a["chave"]]
        # O nome da linha é o da Controladoria, e fica. Quando o embrião "LINDEZA
        # DA PAO GRANDE X FIGO DO YURI - 26/03/2025 RECEP 495" nasce e vira
        # "PAQUITA DA PAO GRANDE", ela não renomeia a linha — e reescrever
        # criaria diferença em cima de um animal que não se moveu. A exceção são
        # as receptoras, cujo nome É a contagem do rebanho ("RECEPTORAS 120") e
        # que ela atualiza todo mês.
        atual = ws.cell(row=r, column=col["nome"]).value
        if not atual or RX_RECEPTORAS.match(norm(atual)):
            ws.cell(row=r, column=col["nome"]).value = a["nome"]
        ws.cell(row=r, column=col["sufixo"]).value = a["sufixo"]
        ws.cell(row=r, column=col["categoria"]).value = a["categoria"]
        ws.cell(row=r, column=col["status"]).value = a["status"]
        ws.cell(row=r, column=col["cota"]).value = a["cota"]
        ws.cell(row=r, column=col["base"]).value = a["base"] or None
        for campo in ("compras", "embrioes", "venda", "morte_doacao", "reavaliacao"):
            ws.cell(row=r, column=col[campo]).value = a[campo] or None
        partes = "+".join(
            f"{get_column_letter(col[k])}{r}"
            for k in ("base", "compras", "embrioes", "venda", "reavaliacao", "morte_doacao")
        )
        ws.cell(row=r, column=col["final"]).value = f"=+{partes}"

    ws.cell(row=lin_cab, column=col["final"]).value = f"VALOR EM {rot_mes}"
    prim, ult = lin_cab + 1, lin_tot - 1
    for campo in ("cota", "base", "compras", "embrioes", "venda",
                  "morte_doacao", "reavaliacao", "final"):
        letra = get_column_letter(col[campo])
        ws.cell(row=lin_tot, column=col[campo]).value = (
            f"=SUBTOTAL(9,{letra}{prim}:{letra}{ult})")
    if ws.auto_filter.ref:
        ws.auto_filter.ref = (
            f"{get_column_letter(col['nome'] - 1)}{lin_cab}"
            f":{get_column_letter(col['final'])}{ult}")

    confrontos(ws, col, ncols, lin_tot, desloca_plantel, dados, rot_mes)
    return len(novos), len(dados["movimentacoes"])


def confrontos(ws, col, ncols, lin_tot, desloca_plantel, dados, rot_mes):
    """Embaixo dos totais a planilha guarda três conferências, cada uma marcada
    por um rótulo à direita. Elas apontam para a linha de totais e para as outras
    abas por endereço fixo, e `insert_rows` do openpyxl não mexe em fórmula — sem
    reescrever, todas passariam a somar a linha errada."""
    ref = {k: f"{get_column_letter(col[k])}{lin_tot}"
           for k in ("base", "compras", "embrioes", "venda", "morte_doacao",
                     "reavaliacao", "final")}
    plantel_tot = 378 + desloca_plantel      # AG378 no template
    formulas = {
        "CONTRONTO DESTA PLANILHA":
            f"=+{ref['base']}+{ref['venda']}+{ref['reavaliacao']}-{ref['final']}"
            f"+{ref['compras']}+{ref['embrioes']}+{ref['morte_doacao']}",
        "CONFRONTO COM PLANTEL": f"=+{ref['final']}-PLANTEL!AG{plantel_tot}",
        "CONFRONTO COM RESUMO": f"=+{ref['final']}-'Resumo Contabil'!O14",
    }
    mes = dados["mes"]
    causas = dados["resumo"].get(mes, {})
    for r in range(lin_tot + 1, lin_tot + 12):
        for c in range(1, ncols + 5):
            rot = norm(ws.cell(row=r, column=c).value)
            if rot in formulas:
                ws.cell(row=r, column=col["final"]).value = formulas[rot]
        # a linha "Movs em <mês>" é a anotação do movimento do mês, digitada à
        # mão no fechamento — aqui ela já sai das causas apuradas
        if norm(ws.cell(row=r, column=col["base"]).value).startswith("MOVS EM"):
            ws.cell(row=r, column=col["base"]).value = f"Movs em {rot_mes}"
            for campo in ("compras", "embrioes", "venda", "morte_doacao", "reavaliacao"):
                ws.cell(row=r, column=col[campo]).value = causas.get(campo) or None


def escreve_resumo(ws, dados: dict):
    """Uma coluna por mês; saldos são fórmula e não se escreve por cima."""
    lin_cab, cols = cabecalho(ws, "Título", ate_linha=10)
    alvo = {
        "(+) COMPRAS": "compras",
        "(+) PRODUCAO EMBRIOES": "embrioes",
        "(-) BAIXA VENDAS": "venda",
        "(-) BAIXA MORTES E DOACOES": "morte_doacao",
        "(+/-) REAVALIACOES": "reavaliacao",
    }
    linhas = {}
    for r in range(lin_cab + 1, min(ws.max_row, lin_cab + 40) + 1):
        for c in range(1, 4):
            rot = norm(ws.cell(row=r, column=c).value)
            for chave in alvo:
                if rot.startswith(chave):
                    linhas.setdefault(chave, r)
    for mes, vals in dados["resumo"].items():
        rot = MES_PT[int(mes.split("-")[1]) - 1]
        c = cols.get(norm(rot))
        if not c:
            continue
        for chave, campo in alvo.items():
            r = linhas.get(chave)
            if r:
                ws.cell(row=r, column=c).value = vals[campo]


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mes", help="mês do fechamento, MM/AAAA")
    ap.add_argument("--template", help="mapa a usar de base (padrão: o do mês anterior)")
    ap.add_argument("--saida", help="arquivo a gerar (padrão: na pasta atual — "
                                    "passe o caminho do Drive para publicar)")
    a = ap.parse_args()

    mm, ano = a.mes.split("/")
    mes = f"{ano}-{int(mm):02d}"
    ant_n, ant_ano = (12, int(ano) - 1) if int(mm) == 1 else (int(mm) - 1, int(ano))
    mes_ant = f"{ant_ano}-{ant_n:02d}"
    rot_mes = f"{MES_PT[int(mm) - 1]}/{ano}"

    template = Path(a.template) if a.template else acha_template(mes_ant)
    # o padrão grava AQUI, não no Drive: o mapa é documento da Controladoria e
    # uma execução de teste não pode aparecer na pasta compartilhada. Para
    # publicar, passa-se --saida com o caminho de lá, de propósito.
    saida = Path(a.saida) if a.saida else (
        Path.cwd() / f"Plantel Haras Pao Grande - Movimentação Jan a Dez {ano} "
                     f"({MES_PT[int(mm) - 1]} {ano}).xlsx")
    if saida.resolve() == template.resolve():
        sys.exit("saída igual ao template: isso sobrescreveria a base. Use --saida.")

    dados = motor(mes)
    print(f"  template: {template.name}", flush=True)
    wb = openpyxl.load_workbook(template)

    esc, novos_pl = escreve_plantel(wb["PLANTEL"], dados)
    print(f"  PLANTEL: {esc} linhas ({novos_pl} novas)", flush=True)
    novos_mov, tot_mov = escreve_movimentacoes(
        wb["Movimentações"], dados, rot_mes, novos_pl)
    print(f"  Movimentações: {tot_mov} linhas ({novos_mov} novas)", flush=True)
    escreve_resumo(wb["Resumo Contabil"], dados)
    print("  Resumo Contabil: colunas de causa preenchidas", flush=True)

    # o título carrega o período do mapa
    ws = wb["Movimentações"]
    for r in range(1, 5):
        for c in range(1, 6):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "MAPA MOVIMENTA" in norm(v):
                ws.cell(row=r, column=c).value = re.sub(r"(A\s+)\S+$", r"\g<1>" + rot_mes, v)

    print(f"  gravando {saida.name} ...", flush=True)
    wb.save(saida)
    print(f"pronto: {saida}")


if __name__ == "__main__":
    main()
