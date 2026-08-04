"""Tira do PPTX de junho/2026 o conteúdo que NÃO sai de planilha.

Comentários do DRE, programação e resultados de exposições, histórico de manejo
e fotos: nada disso tem base de dados por trás — é escrito e fotografado a cada
mês. Em vez de deixar esses slides como placeholder, este script lê o último
deck aprovado e grava o conteúdo em `_docs/comite_conteudo.json`, que passa a ser
o arquivo que a pessoa edita todo mês. O `build_comite.py` lê esse JSON.

Roda UMA vez pra semear o histórico (ou de novo, se um deck novo virar a
referência). O uso normal é editar o JSON direto.

A leitura é por geometria: o gerador do deck original não usou tabelas de
verdade, só caixas de texto posicionadas, então cada slide é reconstruído
agrupando as caixas por linha (mesmo `top`) e ordenando por `left`.

Uso: python hub/tools/extrair_conteudo.py
"""
import json
import re
import sys
from pathlib import Path

HUB = Path(__file__).resolve().parent.parent
REPO = HUB.parent
PPTX = REPO / "ComitêHPG" / "RELATORIO MENSAL_PG_JUNHO26.pptx"
OUT = REPO / "_docs" / "comite_conteudo.json"
FOTOS = HUB / "assets" / "comite" / "fotos"
MES = "2026-06"                      # mês do deck de referência

EMU_IN = 914400
LARG_FOTO = 1400                     # px — foto de slide não precisa de mais


def caixas(slide):
    """[(top_in, left_in, texto)] das caixas com texto, em ordem de leitura."""
    out = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t:
            out.append(((sh.top or 0) / EMU_IN, (sh.left or 0) / EMU_IN, t))
    return sorted(out, key=lambda x: (round(x[0], 1), x[1]))


def linhas(slide, tol=0.12, desde=0.9):
    """Agrupa as caixas em linhas: mesma altura (± tol) = mesma linha."""
    grupos = []
    for top, left, txt in caixas(slide):
        if top < desde:
            continue
        if grupos and abs(grupos[-1][0] - top) <= tol:
            grupos[-1][1].append((left, txt))
        else:
            grupos.append((top, [(left, txt)]))
    return [[t for _, t in sorted(g, key=lambda x: x[0])] for _, g in grupos]


# ------------------------------------------------------------------ S08
def comentarios(slide):
    """3 caixas por linha: categoria · texto · delta. O delta às vezes cai na
    linha de baixo (a caixa é mais alta), então casa pelo formato R$."""
    itens, pend = [], None
    for top, left, txt in caixas(slide):
        if top < 0.9:
            continue
        if re.fullmatch(r"[+\-−]?R\$ ?[\d.,]+k?", txt):
            if pend:
                pend["delta"] = txt
                itens.append(pend)
                pend = None
            continue
        if left < 1.5:
            if pend:
                itens.append(pend)
            pend = {"cat": txt, "txt": "", "delta": ""}
        elif pend:
            pend["txt"] = txt
    if pend:
        itens.append(pend)
    return [i for i in itens if i["txt"]]


# ------------------------------------------------------------------ S23
def programacao(slide):
    ls = linhas(slide, desde=1.0)
    if not ls:
        return []
    return [l for l in ls[1:] if len(l) >= 3]      # ls[0] é o cabeçalho


# ------------------------------------------------------------ S24–S27
def resultados(slide):
    """Duas colunas de blocos: NOME DO ANIMAL em caixa alta, prêmios com 🏆."""
    cx = caixas(slide)
    titulo = cx[0][2] if cx else ""
    sub = next((t for _, l, t in cx if l < 1.0 and t != titulo), "")
    animais, atual = [], None
    for top, left, txt in sorted(cx, key=lambda x: (x[1] > 4.5, x[0])):
        if top < 0.9:
            continue
        if txt.startswith("🏆"):
            if atual:
                atual["premios"].append(txt.lstrip("🏆 ").strip())
            continue
        if txt == sub:
            continue
        atual = {"nome": txt, "premios": []}
        animais.append(atual)
    return {"titulo": titulo, "sub": sub, "animais": [a for a in animais if a["premios"]]}


# ------------------------------------------------------------------ S38
def manejo(slide):
    """Mês curto numa caixa à esquerda, texto do lado.

    Os dois não têm o mesmo `top` — a caixa do mês fica alguns centésimos ABAIXO
    da do texto. Percorrer em ordem de leitura casava o texto com o mês seguinte
    (Jan recebia o texto de Fev) e perdia o primeiro. Por isso as duas colunas
    são coletadas separadas e emparelhadas por posição.
    """
    meses = [(t, x) for t, l, x in caixas(slide) if t >= 0.9 and l < 1.0 and len(x) <= 4]
    textos = [(t, x) for t, l, x in caixas(slide) if t >= 0.9 and l >= 1.0]
    meses.sort(); textos.sort()
    return [[m, txt] for (_, m), (_, txt) in zip(meses, textos)]


# --------------------------------------------------------------- fotos
def fotos(prs, de, ate):
    """Salva as imagens dos slides de foto, redimensionadas. O PPTX tem 26 MB
    quase todos em foto; no deck web isso não pode ir junto do jeito que está."""
    from PIL import Image
    import io
    FOTOS.mkdir(parents=True, exist_ok=True)
    for f in FOTOS.glob("*.jpg"):
        f.unlink()
    saidas, n = [], 0
    for i, slide in enumerate(prs.slides, 1):
        if not (de <= i <= ate):
            continue
        for sh in slide.shapes:
            if sh.shape_type is None or "PICTURE" not in str(sh.shape_type):
                continue
            if (sh.width or 0) / EMU_IN < 1.2:      # logo do cabeçalho
                continue
            try:
                im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
            except Exception:
                continue
            if im.width > LARG_FOTO:
                im = im.resize((LARG_FOTO, round(im.height * LARG_FOTO / im.width)), Image.LANCZOS)
            n += 1
            nome = f"foto{n:02d}.jpg"
            im.save(FOTOS / nome, "JPEG", quality=82, optimize=True)
            saidas.append(f"fotos/{nome}")
    return saidas


def run():
    if not PPTX.exists():
        sys.exit(f"deck de referência não encontrado: {PPTX}")
    from pptx import Presentation
    prs = Presentation(PPTX)
    s = list(prs.slides)

    conteudo = {
        "_leia": ("Conteúdo do comitê que NÃO sai de planilha. Semeado do deck de "
                  "junho/2026 por hub/tools/extrair_conteudo.py; daqui pra frente é "
                  "editado à mão. Cada mês novo é uma chave 'AAAA-MM'; o build usa o "
                  "mês pedido ou, se não existir, o mês anterior mais recente."),
        MES: {
            "comentarios": comentarios(s[7]),
            "exposicoes": {
                "programacao": programacao(s[22]),
                "resultados": [resultados(s[i]) for i in (23, 24, 25, 26)],
            },
            "manejo": manejo(s[37]),
            "fotos": fotos(prs, 39, 51),
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(conteudo, ensure_ascii=False, indent=2), encoding="utf-8")
    c = conteudo[MES]
    print(f"[conteudo] {len(c['comentarios'])} comentários · "
          f"{len(c['exposicoes']['programacao'])} eventos · "
          f"{len(c['exposicoes']['resultados'])} exposições com resultado · "
          f"{len(c['manejo'])} meses de manejo · {len(c['fotos'])} fotos")
    print(f"           -> {OUT.relative_to(REPO)}  (fotos em {FOTOS.relative_to(REPO)})")


if __name__ == "__main__":
    run()
